"""Crude assay extraction: raw assay spreadsheets -> one table of crude-distillation streams.

    python -m nirnay.cases._assays        # rewrites data/cases/raw/refinery/crude_streams.csv

The refinery case studies read only the CSV, so they need no spreadsheet library; this module
(which needs openpyxl and xlrd) documents and reproduces how every number in the CSV was obtained.

The CSV holds the public-domain SPR crudes only. ExxonMobil's terms of use allow redistribution of
its assay files only complete and unaltered, so their numbers are not written to the CSV: they are
extracted at run time from the unaltered files when a case is built with include_exxonmobil=True.

Sources
  * US DOE Strategic Petroleum Reserve crude assays (8 streams). Works of the US Government, public
    domain (17 U.S.C. 105). Cut columns: Gas (C2-C4), fractions 1-7 with end points 175, 250, 375,
    530, 650, 850, 1050 deg F, residuum 1050 F+.
  * ExxonMobil crude assays (5 crudes: Upper Zakum, Azeri BTC, CPC Blend, Erha, Qua Iboe). Provided by
    ExxonMobil Technology & Engineering Company "for information", no warranty; use is subject to
    exxonmobil.com terms of use. Cut columns (deg C): C4-, C5-65, 65-100, 100-150, 150-200, 200-250,
    250-300, 300-350, 350-370, 370-450, 450-500, 500-550, 550+.

Model streams and the assay cuts aggregated into them
                   SPR (deg C, converted from deg F)       ExxonMobil (deg C)
  LPG              C2-C4                                   IBP-C4
  LN  light naph.  C5-79                                   C5-100
  HN  heavy naph.  79-191                                  100-150
  KERO             191-277                                 150-250
  GO  gas oil      277-343                                 250-370
  VGO              343-566                                 370-550
  VR               566+                                    550+

Aggregation of sub-cuts: mass yield and volume yield add; density = total mass / total volume;
sulphur is mass-weighted; RON, cetane index, smoke point and naphthalenes are volume-weighted (the linear
blending approximation also used in the LP) and left empty unless every sub-cut reports a value.
SPR densities are relative densities 60/60 F; they are multiplied by the density of water at 60 F,
0.999016 g/cm3, to give g/cm3 at 15.6 C.
"""
from __future__ import annotations

import csv
from pathlib import Path

from ._builder import RAW

DIR = RAW / "refinery"
OUT = DIR / "crude_streams.csv"
STREAMS = ["LPG", "LN", "HN", "KERO", "GO", "VGO", "VR"]     # plus a "CRUDE" row per crude
FIELDS = ["crude", "source", "file", "stream", "cut", "wt_pct", "vol_pct", "density", "sulphur_wt_pct",
          "ron", "cetane_index", "smoke_point_mm", "naphthalenes_vol_pct"]

SPR_FILES = {
    "Bayou Choctaw Sour": "spr/BayouChoctawSrAssay.xls",
    "Bayou Choctaw Sweet": "spr/DOE_(80399)_BayouChoctawSweetComprehensiveAssay.xls",
    "Big Hill Sour": "spr/BigHillSrAssay.xlsx",
    "Big Hill Sweet": "spr/BigHillSwAssay.xlsx",
    "Bryan Mound Sour": "spr/BM_SourCalAssay.xlsx",
    "Bryan Mound Sweet": "spr/BryanMoundSwAssay.xlsx",
    "West Hackberry Sour": "spr/WestHackberrySrAssay.xlsx",
    "West Hackberry Sweet": "spr/WestHackberrySwAssay.xlsx",
}
EM_FILES = {
    "Upper Zakum": "exxonmobil/upper_zakum.xlsx",
    "Azeri BTC": "exxonmobil/azeri_btc.xlsx",
    "CPC Blend": "exxonmobil/cpc_blend.xlsx",
    "Erha": "exxonmobil/erha.xlsx",
    "Qua Iboe": "exxonmobil/qua_iboe.xlsx",
}
# column offsets (0 = first data column after the label column)
SPR_CUTS = {"LPG": ([0], "C2-C4"), "LN": ([1], "C5-79C"), "HN": ([2, 3], "79-191C"),
            "KERO": ([4], "191-277C"), "GO": ([5], "277-343C"), "VGO": ([6, 7], "343-566C"),
            "VR": ([10], "566C+")}
EM_CUTS = {"LPG": ([("IBP", "C4")], "IBP-C4"), "LN": ([("C5", "65"), ("65", "100")], "C5-100C"),
           "HN": ([("100", "150")], "100-150C"),
           "KERO": ([("150", "200"), ("200", "250")], "150-250C"),
           "GO": ([("250", "300"), ("300", "350"), ("350", "370")], "250-370C"),
           "VGO": ([("370", "450"), ("450", "500"), ("500", "550")], "370-550C"),
           "VR": ([("550", "FBP")], "550C+")}
WATER_60F = 0.999016


def _num(v):
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip().lstrip("*")
        try:
            return float(s)
        except ValueError:
            return None
    return None


def _aggregate(parts):
    """parts: list of dicts wt, vol, dens, S, ron, ci, sp for sub-cuts."""
    wt = sum(p["wt"] for p in parts)
    vol = sum(p["vol"] for p in parts) if all(p["vol"] is not None for p in parts) else None
    out = {"wt_pct": wt, "vol_pct": vol}
    if all(p["dens"] for p in parts):
        out["density"] = wt / sum(p["wt"] / p["dens"] for p in parts)
    else:
        out["density"] = None
    out["sulphur_wt_pct"] = (sum(p["wt"] * p["S"] for p in parts) / wt
                             if all(p["S"] is not None for p in parts) and wt else None)
    for key, name in (("ron", "ron"), ("ci", "cetane_index"), ("sp", "smoke_point_mm"),
                      ("nap", "naphthalenes_vol_pct")):
        out[name] = (sum(p["vol"] * p[key] for p in parts) / vol
                     if all(p[key] is not None for p in parts) and vol else None)
    return out


def _spr(path):
    if path.suffix == ".xlsx":
        import openpyxl
        ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        import xlrd
        s = xlrd.open_workbook(path).sheet_by_index(0)
        rows = [[s.cell_value(i, j) for j in range(s.ncols)] for i in range(s.nrows)]
    lab = {}
    seen_fraction = False
    for r in rows:
        name = r[1].strip() if len(r) > 1 and isinstance(r[1], str) else None
        if name == "Fraction":
            seen_fraction = True
        if seen_fraction and name and name not in lab:
            lab[name] = r[2:14]
    return lab


def _spr_head(path):
    """Whole-crude relative density and sulphur (the block above the cut table)."""
    if path.suffix == ".xlsx":
        import openpyxl
        rows = [list(r) for r in openpyxl.load_workbook(path, data_only=True).worksheets[0].iter_rows(values_only=True)]
    else:
        import xlrd
        sh = xlrd.open_workbook(path).sheet_by_index(0)
        rows = [[sh.cell_value(i, j) for j in range(sh.ncols)] for i in range(sh.nrows)]
    head = {}
    for r in rows:
        name = r[1].strip() if len(r) > 1 and isinstance(r[1], str) else None
        if name == "Fraction":
            break
        if name in ("Relative Density, 60/60° F", "Sulfur, mass %") and name not in head:
            head[name] = next(_num(v) for v in r[2:] if _num(v) is not None)
    return head


def _em(path):
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    lab = {}
    mol = {}
    for row in ws.iter_rows():
        cells = [c for c in row if hasattr(c, "column_letter")]
        name = row[1].value if len(row) > 1 else None
        if isinstance(name, str) and name.strip() and name.strip() not in lab:
            lab[name.strip()] = {c.column_letter: c.value for c in cells}
        # molecule table: a text cell, then its value in the next non-empty cell of the row
        filled = [c for c in cells if c.value not in (None, "", " ")]
        for a, b in zip(filled, filled[1:]):
            if isinstance(a.value, str) and isinstance(b.value, (int, float)):
                mol.setdefault(a.value.strip(), b.value)
    return lab, mol


def _em_columns(lab):
    """{(start, end): column letter} from the Start/End header rows (the 370-FBP atmospheric
    residue column is skipped because 370-450 also starts at 370)."""
    start = next(v for k, v in lab.items() if k.startswith("Start"))
    end = next(v for k, v in lab.items() if k.startswith("End"))
    cols = {}
    for col, a in start.items():
        b = end.get(col)
        if a in (None, " ", "") or b in (None, " ", "") or col in ("A", "B", "C"):
            continue
        key = (str(a).strip(), str(b).strip())
        if key == ("370", "FBP"):
            continue
        cols.setdefault(key, col)
    return cols


def extract() -> list[dict]:
    out = []
    for crude, rel in SPR_FILES.items():
        lab = _spr(DIR / rel)
        get = lambda key, k: _num(lab[key][k]) if key in lab and k < len(lab[key]) else None
        for s, (cols, cut) in SPR_CUTS.items():
            parts = [{"wt": get("mass %", k), "vol": get("Vol. %", k),
                      "dens": (get("Relative Density, 60/60° F", k) or 0) * WATER_60F or None,
                      "S": get("Sulfur, mass %", k), "ron": get("Research Octane Number", k),
                      "ci": get("Cetane Index", k), "sp": get("Smoke point, mm", k),
                      "nap": get("Naphthalenes, Vol. %", k)} for k in cols]
            out.append({"crude": crude, "source": "US DOE SPR", "file": rel, "stream": s, "cut": cut,
                        **_aggregate(parts)})
        head = _spr_head(DIR / rel)
        out.append({"crude": crude, "source": "US DOE SPR", "file": rel, "stream": "CRUDE",
                    "cut": "whole crude", "wt_pct": 100.0, "vol_pct": 100.0,
                    "density": head["Relative Density, 60/60° F"] * WATER_60F,
                    "sulphur_wt_pct": head["Sulfur, mass %"], "ron": None, "cetane_index": None,
                    "smoke_point_mm": None, "naphthalenes_vol_pct": None})
    for crude, rel in EM_FILES.items():
        lab, mol = _em(DIR / rel)
        cols = _em_columns(lab)
        get = lambda key, col: _num(lab[key].get(col)) if key in lab else None
        dens_key = next(k for k in lab if k.startswith("Density @ 15"))
        for s, (cuts, cut) in EM_CUTS.items():
            if s == "LPG" and ("IBP", "C4") not in cols:
                # no IBP-C4 column: C1-C4 from the light-hydrocarbon table (wt% on crude)
                wt = sum(float(mol[k]) for k in ("methane + ethane", "propane", "isobutane", "n-butane"))
                out.append({"crude": crude, "source": "ExxonMobil", "file": rel, "stream": s,
                            "cut": "C1-C4 (light hydrocarbon table)", "wt_pct": wt, "vol_pct": None,
                            "density": None, "sulphur_wt_pct": None, "ron": None,
                            "cetane_index": None, "smoke_point_mm": None, "naphthalenes_vol_pct": None})
                continue
            parts = [{"wt": get("Yield (% wt)", cols[c]), "vol": get("Yield (% vol)", cols[c]),
                      "dens": get(dens_key, cols[c]), "S": get("Total Sulfur (% wt)", cols[c]),
                      "ron": get("RON (Clear)", cols[c]), "ci": get("Cetane Index (D4737A)", cols[c]),
                      "sp": get("Smoke Point (mm)", cols[c]),
                      "nap": get("Naphthalenes (% vol)", cols[c])} for c in cuts]
            out.append({"crude": crude, "source": "ExxonMobil", "file": rel, "stream": s, "cut": cut,
                        **_aggregate(parts)})
        out.append({"crude": crude, "source": "ExxonMobil", "file": rel, "stream": "CRUDE",
                    "cut": "whole crude", "wt_pct": 100.0, "vol_pct": 100.0,
                    "density": get(dens_key, "C"), "sulphur_wt_pct": get("Total Sulfur (% wt)", "C"),
                    "ron": None, "cetane_index": None, "smoke_point_mm": None,
                    "naphthalenes_vol_pct": None})
    return out


def write(rows=None) -> Path:
    rows = rows if rows is not None else [r for r in extract() if r["source"] == "US DOE SPR"]
    with open(OUT, "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=FIELDS)
        w.writeheader()
        for r in rows:
            w.writerow({k: ("" if r.get(k) is None else (f"{r[k]:.6g}" if isinstance(r[k], float) else r[k]))
                        for k in FIELDS})
    return OUT


def load(include_exxonmobil: bool = False) -> dict:
    """{crude: {stream: {field: value}}}: the SPR crudes from the CSV, and optionally the five
    ExxonMobil crudes extracted from the unaltered assay files (needs openpyxl)."""
    data: dict = {}
    with open(OUT, newline="") as fh:
        rows = list(csv.DictReader(fh))
    if include_exxonmobil:
        rows += [{k: ("" if r.get(k) is None else r[k]) for k in FIELDS}
                 for r in extract() if r["source"] == "ExxonMobil"]
    for r in rows:
        rec = {k: (float(v) if v not in ("", None) and k not in ("crude", "source", "file", "stream", "cut")
                   else (v if v != "" else None)) for k, v in r.items()}
        data.setdefault(r["crude"], {})[r["stream"]] = rec
    return data


if __name__ == "__main__":
    p = write()
    print("wrote", p)
